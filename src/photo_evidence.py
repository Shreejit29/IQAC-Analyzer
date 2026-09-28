from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from typing import Any

import fitz  # PyMuPDF
import numpy as np
from PIL import Image, ImageChops, ImageStat
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

try:
    import cv2
except Exception:  # pragma: no cover - handled by a graceful fallback
    cv2 = None


NOT_IDENTIFIED = "Not Identified"


@dataclass
class ExtractedPhoto:
    image_bytes: bytes
    image_ext: str
    report_name: str
    source_page: int
    activity_title: str
    activity_date: str
    activity_time: str
    venue: str
    extraction_method: str = "Embedded image"


_CLOCK = r"(?:[0-9]{1,2}[.:][0-9]{2}|[0-9]{1,2})\s*(?:a\.m\.?|p\.m\.?|am|pm)"
_TIME_PATTERNS = [
    re.compile(rf"\b(?:time|timing)\s*[:\-]?\s*({_CLOCK}(?:\s*(?:to|[-–]|onwards)\s*{_CLOCK})?)\b", re.I),
    re.compile(rf"(?<![0-9.])\b({_CLOCK})\b", re.I),
]


def _clean(value: Any) -> str:
    if value is None:
        return NOT_IDENTIFIED
    text = " ".join(str(value).split()).strip()
    return text or NOT_IDENTIFIED


def _page_range(value: str) -> list[int]:
    text = _clean(value)
    if text == NOT_IDENTIFIED:
        return []
    nums = [int(x) for x in re.findall(r"\d+", text)]
    if not nums:
        return []
    if len(nums) == 1:
        return nums
    start, end = nums[0], nums[1]
    if 1 <= start <= end <= 10000 and end - start <= 100:
        return list(range(start, end + 1))
    return [start]


def _activity_matches_page(record: dict[str, Any], page_no: int, page_text: str) -> tuple[int, int]:
    score = 0
    distance = 9999
    source_pages = _page_range(_clean(record.get("Source Page")))
    if source_pages:
        if page_no in source_pages:
            score += 8
            distance = 0
        else:
            distance = min(abs(page_no - p) for p in source_pages)
            if distance <= 2:
                score += max(1, 5 - distance)
    title = _clean(record.get("Activity Title"))
    venue = _clean(record.get("Venue"))
    date = _clean(record.get("Activity Date"))
    haystack = page_text.lower()
    for value, weight in ((title, 7), (venue, 3), (date, 2)):
        if value != NOT_IDENTIFIED and len(value) >= 4:
            key = value.lower()
            if key in haystack:
                score += weight
            else:
                words = [w for w in re.findall(r"[a-zA-Z0-9]+", key) if len(w) >= 5]
                hits = sum(1 for word in words if word in haystack)
                if hits >= max(1, min(3, len(words) // 2)):
                    score += weight // 2
    return score, distance


def _best_record(records: list[dict[str, Any]], report_name: str, page_no: int, page_text: str) -> dict[str, Any] | None:
    candidates = [r for r in records if _clean(r.get("Source Report")) == _clean(report_name)]
    if not candidates:
        return None
    if len(candidates) == 1:
        return candidates[0]
    scored = [(r, *_activity_matches_page(r, page_no, page_text)) for r in candidates]
    scored.sort(key=lambda x: (-x[1], x[2]))
    return scored[0][0] if scored and (scored[0][1] > 0 or scored[0][2] <= 2) else None


def _page_time(page_texts: list[str], preferred_pages: list[int]) -> str:
    search_order: list[int] = []
    for p in preferred_pages:
        if 1 <= p <= len(page_texts) and p not in search_order:
            search_order.append(p)
    for p in preferred_pages:
        for delta in (1, -1, 2, -2, 3, -3):
            q = p + delta
            if 1 <= q <= len(page_texts) and q not in search_order:
                search_order.append(q)
    search_order.extend(p for p in range(1, len(page_texts) + 1) if p not in search_order)

    for p in search_order:
        text = page_texts[p - 1]
        for pattern in _TIME_PATTERNS:
            match = pattern.search(text)
            if match:
                return _clean(match.group(1))
    return NOT_IDENTIFIED


def _color_metrics(image: Image.Image) -> tuple[float, float, float]:
    rgb = image.convert("RGB")
    rgb.thumbnail((180, 180))
    arr = np.asarray(rgb, dtype=np.uint8)
    chroma = arr.max(axis=2).astype(np.int16) - arr.min(axis=2).astype(np.int16)
    color_fraction = float((chroma >= 16).mean())
    channel_mean = arr.mean(axis=(0, 1))
    mean_spread = float(channel_mean.max() - channel_mean.min())
    # Saturation is calculated approximately from RGB chroma.
    saturation_mean = float(chroma.mean())
    return color_fraction, mean_spread, saturation_mean


def _is_color_photo(image_bytes: bytes, min_dimension: int = 260, color_threshold: float = 8.0) -> bool:
    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            width, height = image.size
            if min(width, height) < min_dimension:
                return False
            color_fraction, mean_spread, saturation_mean = _color_metrics(image)
            if mean_spread >= color_threshold and color_fraction >= 0.035:
                return True
            return color_fraction >= 0.12 and saturation_mean >= 18.0
    except Exception:
        return False


def _trim_white_border(image: Image.Image, threshold: int = 247, padding: int = 18) -> Image.Image:
    rgb = image.convert("RGB")
    arr = np.asarray(rgb)
    mask = np.any(arr < threshold, axis=2)
    if not np.any(mask):
        return rgb
    ys, xs = np.where(mask)
    x0, x1 = max(0, int(xs.min()) - padding), min(rgb.width, int(xs.max()) + padding + 1)
    y0, y1 = max(0, int(ys.min()) - padding), min(rgb.height, int(ys.max()) + padding + 1)
    return rgb.crop((x0, y0, x1, y1))


def _pil_jpeg(image: Image.Image, quality: int = 94) -> bytes:
    image = image.convert("RGB")
    output = io.BytesIO()
    image.save(output, format="JPEG", quality=quality, optimize=True)
    return output.getvalue()


def _normalized_hash(image_bytes: bytes) -> str:
    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            image = image.convert("RGB")
            image.thumbnail((64, 64))
            buf = io.BytesIO()
            image.save(buf, format="JPEG", quality=80)
            return hashlib.sha1(buf.getvalue()).hexdigest()
    except Exception:
        return hashlib.sha1(image_bytes).hexdigest()


def _extract_embedded_images(page: fitz.Page) -> list[tuple[bytes, str, str]]:
    found: list[tuple[bytes, str, str]] = []
    seen_xrefs: set[int] = set()
    page_area = max(page.rect.width * page.rect.height, 1.0)
    for img in page.get_images(full=True):
        xref = int(img[0])
        if xref in seen_xrefs:
            continue
        seen_xrefs.add(xref)
        try:
            info = page.parent.extract_image(xref)
            image_bytes = info.get("image")
            ext = str(info.get("ext") or "jpg").lower()
            if not image_bytes or not _is_color_photo(image_bytes):
                continue
            rects = page.get_image_rects(xref, transform=False)
            largest = max(((r.width * r.height) / page_area for r in rects), default=0.0)
            # A very large image can be a scanned page. It is retained only when
            # it is genuinely colorful; otherwise forms/attendance sheets are ignored.
            if largest > 0.85 and not _is_color_photo(image_bytes, min_dimension=400, color_threshold=10.0):
                continue
            with Image.open(io.BytesIO(image_bytes)) as image:
                width, height = image.size
                aspect = width / max(height, 1)
                if largest < 0.85 and (aspect < 0.35 or aspect > 3.2):
                    continue
            found.append((image_bytes, ext, "Embedded image"))
        except Exception:
            continue
    return found


def _find_scanned_color_regions(page: fitz.Page) -> list[tuple[bytes, str, str]]:
    """Find photographic regions embedded inside a scanned page image.

    This is the key addition for scanned PDFs: the entire page is often a
    single JPEG/PNG, so get_images() cannot separate the photographs. We render
    the page, locate sizeable color-rich connected regions, crop them, and use
    those crops as the photo evidence images.
    """
    if cv2 is None:
        return []

    try:
        matrix = fitz.Matrix(1.45, 1.45)
        pix = page.get_pixmap(matrix=matrix, alpha=False, colorspace=fitz.csRGB)
        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        image_array = np.asarray(image)
        hsv = cv2.cvtColor(image_array, cv2.COLOR_RGB2HSV)
        h, w = hsv.shape[:2]
        page_area = float(w * h)

        # Saturation mask: photographs generally contain far more chroma than
        # the white/gray paper around them. Two morphology passes help join
        # nearby colorful regions within the same photograph.
        mask = ((hsv[:, :, 1] >= 45) & (hsv[:, :, 2] >= 55)).astype(np.uint8) * 255
        kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (19, 19))
        mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, kernel, iterations=2)
        mask = cv2.morphologyEx(mask, cv2.MORPH_OPEN, np.ones((5, 5), np.uint8), iterations=1)

        contours, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        boxes: list[tuple[int, int, int, int]] = []
        for contour in contours:
            x, y, bw, bh = cv2.boundingRect(contour)
            area = float(bw * bh)
            if area < page_area * 0.028 or area > page_area * 0.78:
                continue
            if bw < 180 or bh < 160:
                continue
            aspect = bw / max(bh, 1)
            if aspect < 0.40 or aspect > 2.90:
                continue
            pad_x = max(10, int(bw * 0.035))
            pad_y = max(10, int(bh * 0.035))
            x0, y0 = max(0, x - pad_x), max(0, y - pad_y)
            x1, y1 = min(w, x + bw + pad_x), min(h, y + bh + pad_y)
            crop = image.crop((x0, y0, x1, y1))
            color_fraction, _, sat_mean = _color_metrics(crop)
            crop_arr = np.asarray(crop.convert("RGB"))
            crop_gray = cv2.cvtColor(crop_arr, cv2.COLOR_RGB2GRAY)
            edge_density = float((cv2.Canny(crop_gray, 80, 160) > 0).mean())
            laplacian_var = float(cv2.Laplacian(crop_gray, cv2.CV_64F).var())
            # Decorative posters/logos can be colourful but are usually much
            # flatter than photographs. Require visible photographic texture.
            if color_fraction < 0.14 or sat_mean < 16 or edge_density < 0.085 or laplacian_var < 600:
                continue
            boxes.append((x0, y0, x1, y1))

        # Merge heavily overlapping/nearby boxes. This keeps a photo from being
        # split into multiple crops when the scene has separated colorful areas.
        merged: list[tuple[int, int, int, int]] = []
        for box in sorted(boxes, key=lambda b: (b[1], b[0])):
            bx0, by0, bx1, by1 = box
            merged_this = False
            for i, (mx0, my0, mx1, my1) in enumerate(merged):
                gap_x = max(0, max(mx0, bx0) - min(mx1, bx1))
                gap_y = max(0, max(my0, by0) - min(my1, by1))
                overlap_x = max(0, min(mx1, bx1) - max(mx0, bx0))
                overlap_y = max(0, min(my1, by1) - max(my0, by0))
                if (overlap_x > 0 and overlap_y > 0) or (gap_x < 30 and gap_y < 30):
                    merged[i] = (min(mx0, bx0), min(my0, by0), max(mx1, bx1), max(my1, by1))
                    merged_this = True
                    break
            if not merged_this:
                merged.append(box)

        results: list[tuple[bytes, str, str]] = []
        for x0, y0, x1, y1 in merged:
            crop = image.crop((x0, y0, x1, y1)).convert("RGB")
            if min(crop.size) < 180:
                continue
            results.append((_pil_jpeg(crop), "jpg", "Scanned page region"))
        return results
    except Exception:
        return []


def _full_scanned_photo(page: fitz.Page) -> list[tuple[bytes, str, str]]:
    """Use a full-page raster as a photo only when the scan is strongly colorful."""
    try:
        matrix = fitz.Matrix(1.45, 1.45)
        pix = page.get_pixmap(matrix=matrix, alpha=False, colorspace=fitz.csRGB)
        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        color_fraction, mean_spread, sat_mean = _color_metrics(image)
        arr = np.asarray(image.convert("RGB"))
        gray = cv2.cvtColor(arr, cv2.COLOR_RGB2GRAY) if cv2 is not None else None
        edge_density = float((cv2.Canny(gray, 80, 160) > 0).mean()) if cv2 is not None else 0.0
        if color_fraction < 0.18 or sat_mean < 20 or edge_density < 0.08:
            return []
        trimmed = _trim_white_border(image)
        if min(trimmed.size) < 250:
            return []
        return [(_pil_jpeg(trimmed), "jpg", "Scanned full-page photograph")]
    except Exception:
        return []


def extract_color_photos(
    documents: list[tuple[bytes, str]],
    records: list[dict[str, Any]],
) -> list[ExtractedPhoto]:
    """Extract color photographs from PDF reports, including scanned pages."""
    extracted: list[ExtractedPhoto] = []
    global_hashes: set[str] = set()

    for raw, filename in documents:
        if not filename.lower().endswith(".pdf"):
            continue
        try:
            doc = fitz.open(stream=raw, filetype="pdf")
        except Exception:
            continue
        try:
            page_texts = [(page.get_text("text") or "").strip() for page in doc]
            for page_index, page in enumerate(doc, start=1):
                page_text = page_texts[page_index - 1] if page_index <= len(page_texts) else ""
                record = _best_record(records, filename, page_index, page_text)
                preferred_pages = _page_range(_clean(record.get("Source Page"))) if record else [page_index]
                activity_time = (
                    _clean(record.get("Activity Time"))
                    if record and _clean(record.get("Activity Time")) != NOT_IDENTIFIED
                    else _page_time(page_texts, preferred_pages)
                )
                if not record:
                    record = {}

                embedded_candidates = _extract_embedded_images(page)
                page_images = page.get_images(full=True)
                has_full_page_scan = False
                page_area = max(page.rect.width * page.rect.height, 1.0)
                for img in page_images:
                    try:
                        rects = page.get_image_rects(int(img[0]), transform=False)
                        if rects and max((r.width * r.height) / page_area for r in rects) >= 0.85:
                            has_full_page_scan = True
                            break
                    except Exception:
                        continue

                if has_full_page_scan:
                    # A scan usually appears as one full-page raster. In that
                    # case, inspect the rendered page for individual color-photo
                    # regions rather than exporting the whole scanned form.
                    candidates = _find_scanned_color_regions(page)
                    if not candidates:
                        candidates = _full_scanned_photo(page)
                else:
                    # Normal PDF pages already expose their photographs as
                    # separate image objects. Do not re-scan those pages, or
                    # the same photo can be found twice.
                    candidates = embedded_candidates

                for image_bytes, ext, method in candidates:
                    digest = _normalized_hash(image_bytes)
                    if digest in global_hashes:
                        continue
                    global_hashes.add(digest)
                    extracted.append(
                        ExtractedPhoto(
                            image_bytes=image_bytes,
                            image_ext=ext,
                            report_name=filename,
                            source_page=page_index,
                            activity_title=_clean(record.get("Activity Title")),
                            activity_date=_clean(record.get("Activity Date")),
                            activity_time=activity_time,
                            venue=_clean(record.get("Venue")),
                            extraction_method=method,
                        )
                    )
        finally:
            doc.close()

    return extracted


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_text(cell, text: str, bold: bool = False, size: float = 10.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    paragraph.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)


def _set_page_a4(section) -> None:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)


def build_photo_evidence_docx(photos: list[ExtractedPhoto]) -> bytes:
    """Create an editable Word document with one color photo per page."""
    doc = Document()
    _set_page_a4(doc.sections[0])

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    if not photos:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("No qualifying color photographs were detected in the uploaded PDF reports.")
        r.bold = True
        r.font.size = Pt(14)
        return _docx_bytes(doc)

    for idx, photo in enumerate(photos, start=1):
        if idx > 1:
            doc.add_page_break()

        section = doc.sections[-1]
        _set_page_a4(section)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(7)
        run = title.add_run("IQAC — Color Photo Evidence")
        run.bold = True
        run.font.size = Pt(17)
        run.font.name = "Aptos Display"

        table = doc.add_table(rows=4, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(1.55)
        table.columns[1].width = Inches(5.55)
        fields = [
            ("Activity Title", photo.activity_title),
            ("Date", photo.activity_date),
            ("Time", photo.activity_time),
            ("Venue", photo.venue),
        ]
        for row, (label, value) in zip(table.rows, fields):
            row.cells[0].width = Inches(1.55)
            row.cells[1].width = Inches(5.55)
            _set_cell_shading(row.cells[0], "EAF2F8")
            _set_cell_shading(row.cells[1], "F8FAFC")
            _set_cell_text(row.cells[0], label, bold=True, size=9.5)
            _set_cell_text(row.cells[1], _clean(value), size=9.8)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        rr = p.add_run("Photograph")
        rr.bold = True
        rr.font.size = Pt(11)

        with Image.open(io.BytesIO(photo.image_bytes)) as im:
            width, height = im.size
        max_width = 6.75
        max_height = 6.65
        ratio = min(max_width / max(width, 1), max_height / max(height, 1))
        display_width = width * ratio
        display_height = height * ratio

        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(io.BytesIO(photo.image_bytes), width=Inches(display_width), height=Inches(display_height))

        src = doc.add_paragraph()
        src.alignment = WD_ALIGN_PARAGRAPH.CENTER
        src.paragraph_format.space_before = Pt(1)
        src.paragraph_format.space_after = Pt(0)
        rr = src.add_run(
            f"Source Report: {photo.report_name}  |  Source Page: {photo.source_page}  |  "
            f"Extraction: {photo.extraction_method}  |  Photo {idx} of {len(photos)}"
        )
        rr.font.size = Pt(8.5)
        rr.font.color.rgb = None

    return _docx_bytes(doc)


def _docx_bytes(doc: Document) -> bytes:
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
