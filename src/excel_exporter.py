from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .record_utils import COLUMNS, clean

NAVY = "17365D"
BLUE = "2F75B5"
BLUE_LIGHT = "DCEAF7"
PALE_BLUE = "F5F9FD"
GREEN = "E2F0D9"
GREEN_TEXT = "1F6B3A"
RED = "FCE4D6"
RED_TEXT = "9C2F2F"
GREY = "667085"
WHITE = "FFFFFF"
BORDER_COLOUR = "D6DEE8"

THIN = Side(style="thin", color=BORDER_COLOUR)
MEDIUM_NAVY = Side(style="medium", color=NAVY)


def _has_real_value(value: Any) -> bool:
    text = clean(value).strip().lower()
    return text not in {"", "not identified", "none identified", "none", "n/a", "na"}


def build_excel_bytes(records: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "IQAC Master Data"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BLUE

    last_col_num = len(COLUMNS)
    last_col = get_column_letter(last_col_num)

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"].value = "IQAC MASTER ACTIVITY DATA"
    ws["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"].value = (
        f"AI-assisted activity extraction  •  {len(records)} activity record(s)  •  "
        f"Generated {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
    )
    ws["A2"].font = Font(name="Aptos", size=10, italic=True, color=GREY)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 21

    for col in range(1, last_col_num + 1):
        ws.cell(3, col).fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    ws.row_dimensions[3].height = 6

    header_row = 4
    for col_idx, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(header_row, col_idx, name)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=MEDIUM_NAVY)
    ws.row_dimensions[header_row].height = 40

    present_col = COLUMNS.index("Documents Present") + 1
    absent_col = COLUMNS.index("Documents Absent") + 1
    naac_attr_col = COLUMNS.index("NAAC Attribute") + 1
    naac_metric_col = COLUMNS.index("NAAC Metric") + 1

    first_data_row = header_row + 1

    for row_idx, record in enumerate(records, start=first_data_row):
        base_fill = PALE_BLUE if (row_idx - first_data_row) % 2 else WHITE

        for col_idx, column in enumerate(COLUMNS, start=1):
            cell = ws.cell(row_idx, col_idx, clean(record.get(column)))
            cell.font = Font(name="Aptos", size=10, color="1F2937")
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
            cell.fill = PatternFill("solid", fgColor=base_fill)

        present_cell = ws.cell(row_idx, present_col)
        if _has_real_value(record.get("Documents Present")):
            present_cell.fill = PatternFill("solid", fgColor=GREEN)
            present_cell.font = Font(name="Aptos", size=10, color=GREEN_TEXT, bold=True)

        absent_cell = ws.cell(row_idx, absent_col)
        if _has_real_value(record.get("Documents Absent")):
            absent_cell.fill = PatternFill("solid", fgColor=RED)
            absent_cell.font = Font(name="Aptos", size=10, color=RED_TEXT, bold=True)

        for col_idx in (naac_attr_col, naac_metric_col):
            ws.cell(row_idx, col_idx).fill = PatternFill("solid", fgColor=BLUE_LIGHT)
            ws.cell(row_idx, col_idx).font = Font(name="Aptos", size=10, color=NAVY, bold=True)

        longest = max(
            len(clean(record.get("Objective"))),
            len(clean(record.get("Activity Description"))),
            len(clean(record.get("Outcome"))),
            len(clean(record.get("Documents Present"))),
            len(clean(record.get("Documents Absent"))),
        )
        ws.row_dimensions[row_idx].height = min(96, max(32, 30 + longest / 18))

    end_row = first_data_row + len(records) - 1
    ws.auto_filter.ref = f"A{header_row}:{last_col}{end_row if records else header_row}"

    widths = {
        "A": 13, "B": 14, "C": 16, "D": 36, "E": 20, "F": 20,
        "G": 31, "H": 28, "I": 26, "J": 23, "K": 20, "L": 40,
        "M": 44, "N": 40, "O": 31, "P": 31, "Q": 25, "R": 15,
        "S": 48, "T": 48, "U": 34, "V": 14,
    }
    for index in range(1, last_col_num + 1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = widths.get(letter, 22)

    ws.freeze_panes = "D5"
    ws.print_title_rows = "1:4"
    ws.print_area = f"A1:{last_col}{end_row if records else header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.oddFooter.center.text = "IQAC Master Activity Data"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.font = "Aptos"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.font = "Aptos"

    wb.properties.title = "IQAC Master Activity Data"
    wb.properties.subject = "AI-extracted IQAC activity records"
    wb.properties.creator = "IQAC Analyzer"

    output = BytesIO()
    wb.save(output)
    data = output.getvalue()

    check = load_workbook(BytesIO(data), read_only=False, data_only=False)
    check.close()
    return data
