from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _clean(value: Any) -> str:
    if value is None:
        return "Not Identified"
    text = " ".join(str(value).split()).strip()
    return text or "Not Identified"


def build_excel_bytes(documents: list[dict[str, Any]]) -> bytes:
    """Create ONE general extraction worksheet for all document types.

    The Activity Report Analyzer has its own exporter and is not affected by this file.
    Every general-extractor document is represented in the single `General Extractor` sheet.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "General Extractor"

    navy, blue, white, border = "17365D", "2F75B5", "FFFFFF", "D6DEE8"
    thin = Side(style="thin", color=border)
    headers = [
        "Source File", "Document Type", "Document Title", "Document Date",
        "Academic Year", "Department / Organization", "Field", "Extracted Information",
        "Category", "Source Page", "Confidence", "Summary / Context",
    ]

    ws.sheet_view.showGridLines = False
    last = get_column_letter(len(headers))
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = "GENERAL DOCUMENT EXTRACTOR"
    ws["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = (
        f"All document types • {len(documents)} document(s) • "
        f"Generated {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
    )
    ws["A2"].font = Font(name="Aptos", size=10, italic=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="Aptos", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)

    row = 5
    for item in documents:
        result = item["result"]
        common = [
            item["filename"], result.document_type, result.document_title,
            result.document_date, result.academic_year, result.organization_department,
        ]

        # Document-level information is included as normal fields.
        document_fields = [
            ("Document Summary", result.short_summary, "Document", "Not Identified", "High"),
        ]
        if result.key_entities:
            document_fields.append((
                "Key Entities", "; ".join(result.key_entities), "Entities", "Not Identified", "High"
            ))

        extracted = list(document_fields)
        for field in result.fields:
            extracted.append((field.name, field.value, field.category, field.source_page, field.confidence))

        # Keep table information in the SAME sheet instead of creating a Tables sheet.
        # Each table cell becomes an extractable field with its table title and row/column context.
        for table in result.tables:
            table_headers = table.headers or [
                f"Column {i + 1}"
                for i in range(max((len(r) for r in table.rows), default=0))
            ]
            for row_index, values in enumerate(table.rows, 1):
                for col_index, value in enumerate(values):
                    col_name = (
                        table_headers[col_index]
                        if col_index < len(table_headers)
                        else f"Column {col_index + 1}"
                    )
                    extracted.append((
                        f"Table: {table.title} | Row {row_index} | {col_name}",
                        value,
                        "Table",
                        table.source_page,
                        "High",
                    ))

        if not extracted:
            extracted.append(("Information", "Not Identified", "General", "Not Identified", "Low"))

        for field_name, value, category, source_page, confidence in extracted:
            values = common + [
                field_name, value, category, source_page, confidence, result.short_summary
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col, _clean(value))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
            row += 1

    widths = [
        34, 24, 38, 18, 18, 30, 42, 65, 22, 16, 14, 55
    ]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    last_data_row = max(4, row - 1)
    ws.auto_filter.ref = f"A4:{last}{last_data_row}"
    ws.freeze_panes = "A5"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.45

    wb.properties.title = "General Document Extraction"
    wb.properties.creator = "IQAC Analyzer"

    out = BytesIO()
    wb.save(out)
    data = out.getvalue()
    check = load_workbook(BytesIO(data), read_only=False, data_only=False)
    check.close()
    return data
