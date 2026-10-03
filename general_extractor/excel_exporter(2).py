from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _clean(value: Any) -> str:
    if value is None:
        return "Not Identified"
    text = " ".join(str(value).split()).strip()
    return text or "Not Identified"


def build_excel_bytes(documents: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documents"
    fields = wb.create_sheet("Extracted Fields")
    tables = wb.create_sheet("Tables")

    navy, blue, light, white, border = "17365D", "2F75B5", "DCEAF7", "FFFFFF", "D6DEE8"
    thin = Side(style="thin", color=border)

    def setup(sheet, title, headers):
        sheet.sheet_view.showGridLines = False
        last = get_column_letter(len(headers))
        sheet.merge_cells(f"A1:{last}1")
        sheet["A1"] = title
        sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=white)
        sheet["A1"].fill = PatternFill("solid", fgColor=navy)
        sheet["A1"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 32
        sheet.merge_cells(f"A2:{last}2")
        sheet["A2"] = f"General document extraction • {len(documents)} document(s) • Generated {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
        sheet["A2"].font = Font(name="Aptos", size=10, italic=True)
        for i, h in enumerate(headers, 1):
            c = sheet.cell(4, i, h)
            c.fill = PatternFill("solid", fgColor=blue)
            c.font = Font(name="Aptos", size=10, bold=True, color=white)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(bottom=thin)
        sheet.freeze_panes = "A5"

    doc_headers = ["Record ID", "Source File", "Document Type", "Document Title", "Document Date", "Academic Year", "Department / Organization", "Summary", "Key Entities", "Warnings"]
    setup(ws, "GENERAL DOCUMENT INDEX", doc_headers)
    for idx, item in enumerate(documents, 1):
        r = idx + 4
        vals = [
            f"DOC-{idx:04d}", item["filename"], item["result"].document_type, item["result"].document_title,
            item["result"].document_date, item["result"].academic_year, item["result"].organization_department,
            item["result"].short_summary, "; ".join(item["result"].key_entities) or "Not Identified", item.get("warning") or "",
        ]
        for col, value in enumerate(vals, 1):
            c = ws.cell(r, col, _clean(value)); c.alignment = Alignment(vertical="top", wrap_text=True); c.border = Border(bottom=thin)
    widths = [14, 34, 24, 38, 18, 18, 30, 55, 40, 55]
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A4:J{max(4, len(documents)+4)}"

    field_headers = ["Document", "Document Type", "Field", "Value", "Category", "Source Page", "Confidence"]
    setup(fields, "DYNAMIC EXTRACTED FIELDS", field_headers)
    row = 5
    for item in documents:
        for f in item["result"].fields:
            vals = [item["filename"], item["result"].document_type, f.name, f.value, f.category, f.source_page, f.confidence]
            for col, value in enumerate(vals, 1):
                c = fields.cell(row, col, _clean(value)); c.alignment = Alignment(vertical="top", wrap_text=True); c.border = Border(bottom=thin)
            row += 1
    for i, w in enumerate([34, 24, 30, 65, 24, 16, 14], 1): fields.column_dimensions[get_column_letter(i)].width = w
    fields.auto_filter.ref = f"A4:G{max(4, row-1)}"

    table_headers = ["Document", "Document Type", "Table", "Source Page", "Column", "Value", "Row"]
    setup(tables, "EXTRACTED TABLE CONTENT", table_headers)
    row = 5
    for item in documents:
        for table in item["result"].tables:
            headers = table.headers or [f"Column {i+1}" for i in range(max((len(x) for x in table.rows), default=0))]
            for ridx, values in enumerate(table.rows, 1):
                for cidx, value in enumerate(values):
                    vals = [item["filename"], item["result"].document_type, table.title, table.source_page, headers[cidx] if cidx < len(headers) else f"Column {cidx+1}", value, ridx]
                    for col, val in enumerate(vals, 1):
                        c = tables.cell(row, col, _clean(val)); c.alignment = Alignment(vertical="top", wrap_text=True); c.border = Border(bottom=thin)
                    row += 1
    for i, w in enumerate([34, 24, 30, 16, 28, 60, 10], 1): tables.column_dimensions[get_column_letter(i)].width = w
    tables.auto_filter.ref = f"A4:G{max(4, row-1)}"

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.orientation = "landscape"
        sheet.page_margins.left = sheet.page_margins.right = 0.25
        sheet.page_margins.top = sheet.page_margins.bottom = 0.45

    wb.properties.title = "General Document Extraction"
    wb.properties.creator = "IQAC Analyzer"
    out = BytesIO(); wb.save(out); data = out.getvalue()
    check = load_workbook(BytesIO(data), read_only=False, data_only=False); check.close()
    return data
