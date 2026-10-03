from __future__ import annotations

import io
import json
import re
import time
from functools import lru_cache
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field


class GeneralExtractorError(RuntimeError):
    """Raised when general document extraction cannot be completed."""


class ExtractedField(BaseModel):
    name: str = Field(default="Not Identified")
    value: str = Field(default="Not Identified")
    category: str = Field(default="General")
    source_page: str = Field(default="Not Identified")
    confidence: str = Field(default="Medium")


class ExtractedTable(BaseModel):
    title: str = Field(default="Untitled Table")
    headers: list[str] = Field(default_factory=list)
    rows: list[list[str]] = Field(default_factory=list)
    source_page: str = Field(default="Not Identified")


class GeneralDocument(BaseModel):
    document_type: str = Field(default="Other / Unknown")
    document_title: str = Field(default="Not Identified")
    document_date: str = Field(default="Not Identified")
    academic_year: str = Field(default="Not Identified")
    organization_department: str = Field(default="Not Identified")
    short_summary: str = Field(default="Not Identified")
    key_entities: list[str] = Field(default_factory=list)
    fields: list[ExtractedField] = Field(default_factory=list)
    tables: list[ExtractedTable] = Field(default_factory=list)


SYSTEM_PROMPT = """
You are a general institutional-document extraction engine used by a college IQAC office.

Your job is to understand the supplied document and extract information according to the
actual content of that document. This is NOT an activity-report-only extractor.

IMPORTANT RULES
1. First identify what kind of document it is. Examples include activity report, notice,
   circular, committee list, meeting minutes, placement report, certificate, attendance,
   workshop/seminar report, proposal, timetable, achievement record, MoU, research/publication
   document, financial/admin document, or Other / Unknown.
2. Do not force the document into any predefined institutional template.
3. Extract the meaningful fields that actually occur in the document.
4. Field names must describe the document's real content. Examples: "Committee Name",
   "Chairperson", "Company Name", "Students Selected", "Resource Person", "Notice Date",
   "Certificate Holder", "Achievement", "Meeting Decision", etc.
5. Do not invent missing facts. Use exactly "Not Identified" when a value cannot be supported.
6. Preserve names, dates, numbers, titles, reference numbers, organizations and amounts accurately.
7. Consolidate information across the complete document.
8. If tables exist, reproduce their meaningful headers and rows in the tables field.
9. For PDFs, use the PDF page number where possible. For DOCX/TXT, use "Not Identified"
   unless page information is explicitly available.
10. Confidence must be High, Medium, or Low and reflect extraction certainty, not importance.
11. Keep field values concise. Do not copy long paragraphs into fields.
12. key_entities should contain important named people, organizations, companies, committees,
    places or other entities that are clearly present.
13. short_summary should be a concise factual description of what the document contains.
14. This is an extraction system, not a report-writing system and not an accreditation decision engine.
15. Never manufacture NAAC metrics, outcomes, attendance, participants, approvals, or evidence.
""".strip()


@lru_cache(maxsize=4)
def _client(api_key: str):
    try:
        from google import genai
        return genai.Client(api_key=api_key)
    except Exception as exc:
        raise GeneralExtractorError(f"Gemini SDK is not available: {exc}") from exc


def _retryable(exc: Exception) -> bool:
    msg = str(exc).lower()
    return any(x in msg for x in ("429", "500", "502", "503", "504", "timeout", "timed out", "unavailable", "resource exhausted"))


def _extract_local_text(raw: bytes, filename: str) -> tuple[str, bool]:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".txt":
            text = raw.decode("utf-8", errors="replace").strip()
            return text, bool(text)

        if suffix == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(raw))
            chunks: list[str] = []
            for p in doc.paragraphs:
                if p.text.strip():
                    chunks.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        chunks.append(" | ".join(cells))
            text = "\n".join(chunks).strip()
            return text, len(text) >= 30

        if suffix == ".pdf":
            import fitz
            doc = fitz.open(stream=raw, filetype="pdf")
            chunks = []
            for page_no, page in enumerate(doc, start=1):
                text = page.get_text("text").strip()
                if text:
                    chunks.append(f"[Page {page_no}]\n{text}")
            doc.close()
            text = "\n\n".join(chunks).strip()
            return text, len(text) >= 30
    except Exception:
        return "", False
    return "", False


def _generate(client: Any, model: str, contents: Any) -> GeneralDocument:
    try:
        from google.genai import types
        response = client.models.generate_content(
            model=model,
            contents=contents,
            config=types.GenerateContentConfig(
                system_instruction=SYSTEM_PROMPT,
                temperature=0,
                thinking_config=types.ThinkingConfig(thinking_level="minimal"),
                response_mime_type="application/json",
                response_schema=GeneralDocument,
            ),
        )
    except Exception as exc:
        raise GeneralExtractorError(str(exc)) from exc

    parsed = getattr(response, "parsed", None)
    if parsed is not None:
        try:
            return parsed if isinstance(parsed, GeneralDocument) else GeneralDocument.model_validate(parsed)
        except Exception:
            pass

    text = getattr(response, "text", "") or ""
    if not text:
        raise GeneralExtractorError("Gemini returned an empty response.")
    try:
        return GeneralDocument.model_validate_json(text)
    except Exception:
        try:
            return GeneralDocument.model_validate(json.loads(text))
        except Exception as exc:
            raise GeneralExtractorError(f"Gemini returned invalid structured data: {exc}") from exc


def _generate_retry(client: Any, model: str, contents: Any) -> GeneralDocument:
    last: Exception | None = None
    for attempt, delay in enumerate((0, 2, 5), start=1):
        if delay:
            time.sleep(delay)
        try:
            return _generate(client, model, contents)
        except Exception as exc:
            last = exc
            if not _retryable(exc) or attempt == 3:
                break
    raise GeneralExtractorError(str(last) if last else "General extraction failed.")


def _fallback_document(raw: bytes, filename: str) -> GeneralDocument:
    text, usable = _extract_local_text(raw, filename)
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    title = lines[0][:180] if lines else Path(filename).stem
    dates = re.findall(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})\b", text)
    return GeneralDocument(
        document_type="Other / Unknown",
        document_title=title or "Not Identified",
        document_date=dates[0] if dates else "Not Identified",
        short_summary=("Local text extraction completed; AI classification was unavailable." if usable else "Could not extract readable text locally."),
        fields=[
            ExtractedField(name="File Name", value=filename, category="File", confidence="High"),
            ExtractedField(name="Local Text Available", value="Yes" if usable else "No", category="Extraction", confidence="High"),
        ],
    )


def analyze_document(raw: bytes, filename: str, model: str, api_key: str) -> tuple[GeneralDocument, str | None]:
    """Analyze one arbitrary document. Returns (result, warning)."""
    if not api_key.strip():
        return _fallback_document(raw, filename), "GEMINI_API_KEY is not configured; local extraction fallback used."

    client = _client(api_key)
    suffix = Path(filename).suffix.lower()
    prompt = f"Analyze the complete document named '{filename}'. Return only the requested structured JSON."
    try:
        if suffix == ".pdf":
            from google.genai import types
            part = types.Part.from_bytes(data=raw, mime_type="application/pdf")
            result = _generate_retry(client, model, [prompt, part])
        else:
            text, usable = _extract_local_text(raw, filename)
            if not usable:
                raise GeneralExtractorError(f"Could not extract usable text from {filename}.")
            result = _generate_retry(client, model, [prompt, text])
        return result, None
    except Exception as exc:
        fallback = _fallback_document(raw, filename)
        return fallback, f"AI extraction failed: {exc}"
